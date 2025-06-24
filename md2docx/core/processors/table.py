from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

from md2docx.core import line_type
from md2docx.utils import DocStyle


def set_cell_background_color(cell, color: str = 'FFFFFF'):
    """设置表格单元格背景色"""
    try:
        color = color.strip('#')
        shading_elm = parse_xml(
            f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{color}"/>'
        )
        cell._tc.get_or_add_tcPr().append(shading_elm)
    except Exception as e:
        print(f"设置单元格背景色失败: {e}")


def process_table(
        doc: Document,
        lines: list,
        start_idx: int,
        alignment: str = 'center',
    ) -> int:
    """处理表格"""
    table_lines = []
    i = start_idx

    alignment = getattr(WD_TABLE_ALIGNMENT, alignment.upper(), WD_TABLE_ALIGNMENT.CENTER)
    
    # 收集所有表格相关行（包括分隔符）
    table_raw_lines = []
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            break
        if line_type.is_table_row(line) or line_type.is_table_separator(line):
            table_raw_lines.append(line)
            i += 1
        else:
            break
    
    # 过滤掉分隔符行，只保留数据行
    for line in table_raw_lines:
        if not line_type.is_table_separator(line):
            table_lines.append(line)
    
    if len(table_lines) < 1:  # 至少需要一行数据
        return start_idx + 1
    
    # 解析表格数据
    table_data = []
    for line in table_lines:
        # 移除行首尾的|，然后分割
        cells = [cell.strip() for cell in line.strip('|').split('|')]
        table_data.append(cells)
    
    # 创建表格
    if table_data:
        rows = len(table_data)
        cols = len(table_data[0]) if table_data else 0
        
        if rows > 0 and cols > 0:
            table = doc.add_table(rows=rows, cols=cols)
            table.style = 'Table Grid'
            table.alignment = alignment
            table.autofit = True

            # 填充表格数据
            for row_idx, row_data in enumerate(table_data):
                for col_idx, cell_data in enumerate(row_data):
                    if col_idx < cols:  # 确保不超出列数
                        cell = table.cell(row_idx, col_idx)
                        cell.text = cell_data
                        
                        # 设置标题行样式（第一行）
                        if row_idx == 0:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.font.bold = True
    return i
